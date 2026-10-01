"""
xlsx 输入输出模块：题目分布表（1.xlsx）与到发站对表（2.xlsx）的解析、失败回执生成

与参考项目一致（自 server.py 迁出）：
- 分布表布局兼容原模板：「总题目」行 +「存在性」区（单策略/双策略/三策略 ×
  有干扰/无干扰两列）+「选择性」区（评判标准 × 行为约束 × 数量）
- 站对表：每行 [出发站, 到达站]，跳过空行与同站对
"""

from io import BytesIO
from typing import Any, Dict, List, Optional


# ============================================================
# 分布行名称 → 题型参数映射（单策略 / 双策略 / 三策略；多策略一律"一次换乘类混合"）
# ============================================================
_EXISTS_STRATEGY_MAP = {
    "换乘": {"question_type": "transfer", "transfers": 0, "segment_plans": []},
    "买短补长": {"question_type": "short_buy", "transfers": 0, "segment_plans": []},
    "额外购买（前）": {"question_type": "extra_front", "transfers": 0, "segment_plans": []},
    "额外购买（后）": {"question_type": "extra_rear", "transfers": 0, "segment_plans": []},
    # 双策略：前半段直达 + 后半段指定策略
    "买短补长+换乘": {"question_type": "mixed", "transfers": 1, "segment_plans": ["direct", "short_buy"]},
    "额外购买（前）+换乘": {"question_type": "mixed", "transfers": 1, "segment_plans": ["direct", "extra_front"]},
    "额外购买（后）+换乘": {"question_type": "mixed", "transfers": 1, "segment_plans": ["direct", "extra_rear"]},
    # 三策略：前后两段都有策略
    "买短补长+买短补长": {"question_type": "mixed", "transfers": 1, "segment_plans": ["short_buy", "short_buy"]},
    "买短补长+额外购买（前）": {"question_type": "mixed", "transfers": 1, "segment_plans": ["short_buy", "extra_front"]},
    "买短补长+额外购买（后）": {"question_type": "mixed", "transfers": 1, "segment_plans": ["short_buy", "extra_rear"]},
    "额外购买（前）+额外购买（前）": {"question_type": "mixed", "transfers": 1, "segment_plans": ["extra_front", "extra_front"]},
    "额外购买（前）+额外购买（后）": {"question_type": "mixed", "transfers": 1, "segment_plans": ["extra_front", "extra_rear"]},
    "额外购买（后）+额外购买（后）": {"question_type": "mixed", "transfers": 1, "segment_plans": ["extra_rear", "extra_rear"]},
}
_CRITERION_NAME_MAP = {
    "综合考虑": "comprehensive", "最快": "fastest", "最便宜": "cheapest",
    "出发最晚": "depart_latest", "最早到达": "arrive_earliest",
}
_CRITERION_CN = {v: k for k, v in _CRITERION_NAME_MAP.items()}
_BEHAVIOR_NAME_MAP = {
    "不允许换乘": "no_transfer",
    "不允许买短补长与额外购买": "no_short_buy_extra",
}
_BEHAVIOR_CN = {v: k for k, v in _BEHAVIOR_NAME_MAP.items()}
_BEHAVIOR_CN["none"] = "随意"


def criterion_cn(code: str) -> str:
    """评判标准代码 → 中文名"""
    return _CRITERION_CN.get(code, code)


def behavior_cn(code: str) -> str:
    """行为约束代码 → 中文名"""
    return _BEHAVIOR_CN.get(code, code)


def _cell_int(v) -> int:
    """xlsx 单元格 → 非负整数（空值/文本返回 0）"""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return max(0, int(v))
    if isinstance(v, str) and v.strip().isdigit():
        return int(v.strip())
    return 0


def parse_distribution_xlsx(data: bytes) -> Dict[str, Any]:
    """解析 1.xlsx 题目分布表（兼容原模板布局）。

    返回 {total: int|None, exists: [...], selective: [...]}，
    其中 exists 行含 category/name/has_interference/no_interference 及映射后的
    question_type/transfers/segment_plans；selective 行含 criterion/behavior/count。
    """
    import openpyxl
    wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    wb.close()

    result: Dict[str, Any] = {"total": None, "exists": [], "selective": []}
    section: Optional[str] = None  # "exists" / "selective"
    current_category: Optional[str] = None
    current_criterion: Optional[str] = None

    for r in rows:
        cells = [str(c).strip() if c is not None else "" for c in r]
        joined = "".join(cells)
        if "总题目" in joined:
            for c in r:
                if isinstance(c, (int, float)) and not isinstance(c, bool):
                    result["total"] = int(c)
                    break
            continue
        if "存在性" in joined:
            section = "exists"
            current_category = None
            continue
        if "选择性" in joined:
            section = "selective"
            continue
        if not joined:
            continue

        if section == "exists":
            cat = cells[0] if cells[0] in ("单策略", "双策略", "三策略") else current_category
            name = cells[1] if len(cells) > 1 else ""
            config = _EXISTS_STRATEGY_MAP.get(name)
            if config:
                current_category = cat or "单策略"
                result["exists"].append({
                    "category": current_category or "单策略",
                    "name": name,
                    "has_interference": _cell_int(r[2]) if len(r) > 2 else 0,
                    "no_interference": _cell_int(r[3]) if len(r) > 3 else 0,
                    **{k: config[k] for k in ("question_type", "transfers", "segment_plans")},
                })
        elif section == "selective":
            crit = cells[0] if cells else ""
            if crit in _CRITERION_NAME_MAP:
                current_criterion = crit
            # 合并单元格：评判标准列为空时沿用上一行的标准
            if current_criterion is None:
                continue
            beh = cells[1] if len(cells) > 1 else "随意"
            cnt = _cell_int(r[2]) if len(r) > 2 else 0
            result["selective"].append({
                "criterion": _CRITERION_NAME_MAP[current_criterion],
                "behavior": _BEHAVIOR_NAME_MAP.get(beh, "none"),
                "count": cnt,
            })
    return result


def parse_stations_xlsx(data: bytes) -> List[List[str]]:
    """解析 2.xlsx 到发站对：每行 [出发站, 到达站]，跳过空行与同站对。"""
    import openpyxl
    wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    stations: List[List[str]] = []
    for r in rows:
        if not r:
            continue
        a = str(r[0]).strip() if r[0] is not None else ""
        b = str(r[1]).strip() if len(r) > 1 and r[1] is not None else ""
        # 跳过表头（「出发站/到达站」两格或合并单格），否则会被当成一对站白费一次出题尝试
        if "出发站" in a and "到达站" in (a + b):
            continue
        if a and b and a != b:
            stations.append([a, b])
    return stations


def generate_failure_report_xlsx(report: Dict[str, Any]) -> bytes:
    """根据批量出题结果生成同构回执 xlsx 字节流（与 1.xlsx 同构）：
    存在性区每格=该格失败数，选择性区每格=失败数，顶部=总失败数。"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "批量回执"
    ws["A1"] = "总题目"
    ws["B1"] = report.get("total", 0)
    ws["A2"] = "成功"
    ws["B2"] = report.get("success", 0)
    ws["A3"] = "失败（总失败数）"
    ws["B3"] = report.get("failed", 0)

    ws["A5"] = "存在性问题"
    ws["C5"] = "有干扰失败"
    ws["D5"] = "无干扰失败"
    r = 6
    for item in report.get("exists", []):
        ws.cell(row=r, column=1, value=item.get("category", ""))
        ws.cell(row=r, column=2, value=item.get("name", ""))
        ws.cell(row=r, column=3, value=item.get("has_interference_failed", 0))
        ws.cell(row=r, column=4, value=item.get("no_interference_failed", 0))
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="选择性问题")
    ws.cell(row=r, column=3, value="失败数")
    r += 1
    for item in report.get("selective", []):
        ws.cell(row=r, column=1, value=_CRITERION_CN.get(item.get("criterion", ""), item.get("criterion", "")))
        ws.cell(row=r, column=2, value=_BEHAVIOR_CN.get(item.get("behavior", ""), item.get("behavior", "")))
        ws.cell(row=r, column=3, value=item.get("failed", 0))
        r += 1

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
